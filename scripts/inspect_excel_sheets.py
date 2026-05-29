import openpyxl
import pandas as pd
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')

PROJECT_ROOT = Path(r"C:\Users\Ars\projects\university\lab_urfu_2026\false-lv-chords")
raw_path = PROJECT_ROOT / "data" / "raw" / "final_analysis_dataset_Corr_08.05.2026.xlsx"

wb = openpyxl.load_workbook(raw_path, read_only=True)
print("Sheets in raw workbook:", wb.sheetnames)

# Let's inspect the first sheet columns
sheet = wb.active
print("Active sheet name:", sheet.title)

# Read first row
first_row = []
for cell in next(sheet.iter_rows(max_row=1, values_only=True)):
    first_row.append(str(cell))

print(f"Total columns in active sheet: {len(first_row)}")
print("First 50 columns:")
for i, col in enumerate(first_row[:50]):
    print(f"  {i}: {col}")

# Let's search for keywords in all column names of active sheet
keywords = ["давление", "систолическое", "давл", "bp", "sbp", "pressure", "кср", "esd", "тзс", "задн", "posterior", "wall", "thickness", "зс", "тзслжс", "ад"]
matched = []
for i, col in enumerate(first_row):
    col_lower = col.lower()
    if any(k in col_lower for k in keywords):
        matched.append((i, col))

print(f"\nMatched columns in active sheet ({len(matched)}):")
for i, col in matched:
    print(f"  {i}: {col}")
